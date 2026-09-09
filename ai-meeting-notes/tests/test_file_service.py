import io
import tempfile
from pathlib import Path
import pytest
import docx
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from app.services.file_service import FileService


def test_extract_from_txt(tmp_path):
    txt_file = tmp_path / "meeting.txt"
    content = "Alice: Let's discuss the budget.\nBob: Approved."
    txt_file.write_text(content, encoding="utf-8")

    extracted = FileService.extract_text_from_file(txt_file)
    assert "Alice: Let's discuss the budget." in extracted
    assert "Bob: Approved." in extracted


def test_extract_from_docx(tmp_path):
    docx_file = tmp_path / "meeting.docx"
    doc = docx.Document()
    doc.add_paragraph("Quarterly Review Meeting")
    doc.add_paragraph("Decision: Launch next week.")
    doc.save(str(docx_file))

    extracted = FileService.extract_text_from_file(docx_file)
    assert "Quarterly Review Meeting" in extracted
    assert "Decision: Launch next week." in extracted


def test_extract_from_pdf(tmp_path):
    pdf_file = tmp_path / "meeting.pdf"
    doc = SimpleDocTemplate(str(pdf_file))
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Sprint Retrospective Notes", styles["Heading1"]),
        Paragraph("Action item: Alex will fix the database indexing.", styles["Normal"])
    ]
    doc.build(story)

    extracted = FileService.extract_text_from_file(pdf_file)
    assert "Sprint Retrospective Notes" in extracted
    assert "Alex will fix the database indexing" in extracted


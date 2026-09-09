import io
from pathlib import Path
from typing import Optional
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.database import Meeting
from app.utils.logger import logger


class ExportService:
    """Generates professional PDF, DOCX, Markdown, and TXT meeting notes reports."""

    @staticmethod
    def export_to_markdown(meeting: Meeting) -> str:
        """Exports meeting notes to GitHub-flavored markdown."""
        date_str = meeting.date.strftime("%B %d, %Y") if meeting.date else "Not available"
        duration = f"{meeting.duration_minutes:.1f} mins" if meeting.duration_minutes else "Not available"

        md = []
        md.append(f"# {meeting.title}\n")
        md.append(f"**Date:** {date_str}  |  **Duration:** {duration}  |  **Source:** {meeting.source_filename or 'Manual'}\n")
        md.append("---\n")

        # Participants
        if meeting.participants:
            md.append("## 👥 Participants\n")
            for p in meeting.participants:
                role_str = f" ({p.role})" if p.role else ""
                md.append(f"- **{p.name}**{role_str}")
            md.append("\n")

        # Executive Summary
        if meeting.summary:
            md.append("## 📋 Executive Summary\n")
            md.append(f"{meeting.summary.quick_summary}\n\n")

            if meeting.summary.standard_summary:
                md.append("## 📝 Detailed Summary\n")
                md.append(f"{meeting.summary.detailed_summary or meeting.summary.standard_summary}\n\n")

            # Key Points
            if meeting.summary.key_points:
                md.append("## 💡 Key Discussion Points\n")
                for kp in meeting.summary.key_points:
                    md.append(f"- {kp}")
                md.append("\n")

            # Decisions
            if meeting.summary.decisions:
                md.append("## 🎯 Decisions Made\n")
                for d in meeting.summary.decisions:
                    md.append(f"- ✅ {d}")
                md.append("\n")

            # Questions / Open Issues
            if meeting.summary.questions:
                md.append("## ❓ Open Questions & Issues\n")
                for q in meeting.summary.questions:
                    md.append(f"- ❓ {q}")
                md.append("\n")

        # Action Items
        if meeting.action_items:
            md.append("## ⚡ Action Items\n")
            md.append("| Task | Assignee | Deadline | Priority | Status |")
            md.append("| :--- | :--- | :--- | :--- | :--- |")
            for item in meeting.action_items:
                priority = item.priority.value if item.priority else "MEDIUM"
                status = item.status.value if item.status else "PENDING"
                md.append(f"| {item.task} | {item.assignee} | {item.deadline} | **{priority}** | `{status}` |")
            md.append("\n")

        return "\n".join(md)

    @staticmethod
    def export_to_txt(meeting: Meeting) -> str:
        """Exports meeting notes to plain text."""
        date_str = meeting.date.strftime("%B %d, %Y") if meeting.date else "Not available"
        duration = f"{meeting.duration_minutes:.1f} mins" if meeting.duration_minutes else "Not available"

        txt = []
        txt.append("=" * 60)
        txt.append(f"MEETING SUMMARY: {meeting.title.upper()}")
        txt.append(f"Date: {date_str} | Duration: {duration}")
        txt.append(f"Source File: {meeting.source_filename or 'Manual Input'}")
        txt.append("=" * 60 + "\n")

        if meeting.participants:
            txt.append("PARTICIPANTS:")
            for p in meeting.participants:
                role = f" ({p.role})" if p.role else ""
                txt.append(f" - {p.name}{role}")
            txt.append("")

        if meeting.summary:
            txt.append("EXECUTIVE SUMMARY:")
            txt.append(meeting.summary.quick_summary + "\n")

            txt.append("DETAILED SUMMARY:")
            txt.append((meeting.summary.detailed_summary or meeting.summary.standard_summary or "") + "\n")

            if meeting.summary.key_points:
                txt.append("KEY DISCUSSION POINTS:")
                for kp in meeting.summary.key_points:
                    txt.append(f" * {kp}")
                txt.append("")

            if meeting.summary.decisions:
                txt.append("DECISIONS MADE:")
                for d in meeting.summary.decisions:
                    txt.append(f" [DONE] {d}")
                txt.append("")

            if meeting.summary.questions:
                txt.append("OPEN QUESTIONS:")
                for q in meeting.summary.questions:
                    txt.append(f" ? {q}")
                txt.append("")

        if meeting.action_items:
            txt.append("ACTION ITEMS:")
            for idx, item in enumerate(meeting.action_items, 1):
                prio = item.priority.value if item.priority else "MEDIUM"
                stat = item.status.value if item.status else "PENDING"
                txt.append(f" {idx}. [{stat}] {item.task}")
                txt.append(f"    Assignee: {item.assignee} | Deadline: {item.deadline} | Priority: {prio}")
            txt.append("")

        return "\n".join(txt)

    @staticmethod
    def export_to_docx_bytes(meeting: Meeting) -> bytes:
        """Exports meeting notes to Microsoft Word .docx binary stream."""
        doc = docx.Document()

        # Styles
        title_para = doc.add_heading(meeting.title, level=0)
        title_para.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT

        date_str = meeting.date.strftime("%B %d, %Y") if meeting.date else "Not available"
        duration = f"{meeting.duration_minutes:.1f} minutes" if meeting.duration_minutes else "Not available"

        meta_p = doc.add_paragraph()
        meta_p.add_run(f"Date: {date_str}   |   Duration: {duration}   |   Source: {meeting.source_filename or 'Manual'}").italic = True
        doc.add_paragraph()  # spacer

        # Summary
        if meeting.summary:
            doc.add_heading("Executive Summary", level=1)
            doc.add_paragraph(meeting.summary.quick_summary)

            if meeting.summary.detailed_summary:
                doc.add_heading("Detailed Summary", level=1)
                doc.add_paragraph(meeting.summary.detailed_summary)

            if meeting.summary.key_points:
                doc.add_heading("Key Discussion Points", level=1)
                for kp in meeting.summary.key_points:
                    doc.add_paragraph(kp, style="List Bullet")

            if meeting.summary.decisions:
                doc.add_heading("Decisions Made", level=1)
                for d in meeting.summary.decisions:
                    doc.add_paragraph(d, style="List Bullet")

            if meeting.summary.questions:
                doc.add_heading("Open Questions & Issues", level=1)
                for q in meeting.summary.questions:
                    doc.add_paragraph(q, style="List Bullet")

        # Action items table
        if meeting.action_items:
            doc.add_heading("Action Items", level=1)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Task"
            hdr_cells[1].text = "Assignee"
            hdr_cells[2].text = "Deadline"
            hdr_cells[3].text = "Priority"
            hdr_cells[4].text = "Status"

            for item in meeting.action_items:
                row_cells = table.add_row().cells
                row_cells[0].text = item.task
                row_cells[1].text = item.assignee
                row_cells[2].text = item.deadline
                row_cells[3].text = item.priority.value if item.priority else "MEDIUM"
                row_cells[4].text = item.status.value if item.status else "PENDING"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_to_pdf_bytes(meeting: Meeting) -> bytes:
        """Exports meeting notes to a styled PDF binary stream using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )

        styles = getSampleStyleSheet()

        # Custom Styles
        primary_color = colors.HexColor("#1E3A8A")
        text_color = colors.HexColor("#1F2937")

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Title"],
            fontSize=20,
            leading=24,
            textColor=primary_color,
            alignment=0,
            spaceAfter=6,
        )
        meta_style = ParagraphStyle(
            "DocMeta",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#6B7280"),
            spaceAfter=12,
        )
        h1_style = ParagraphStyle(
            "SectionH1",
            parent=styles["Heading1"],
            fontSize=13,
            leading=16,
            textColor=primary_color,
            spaceBefore=12,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyDark",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=text_color,
            spaceAfter=6,
        )
        bullet_style = ParagraphStyle(
            "BulletDark",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=text_color,
            leftIndent=15,
            firstLineIndent=-10,
            spaceAfter=4,
        )

        story = []

        # Title and Header
        story.append(Paragraph(meeting.title, title_style))
        date_str = meeting.date.strftime("%B %d, %Y") if meeting.date else "Not available"
        duration = f"{meeting.duration_minutes:.1f} mins" if meeting.duration_minutes else "Not available"
        meta_text = f"<b>Date:</b> {date_str} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Duration:</b> {duration} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Source:</b> {meeting.source_filename or 'Manual'}"
        story.append(Paragraph(meta_text, meta_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#E5E7EB"), spaceAfter=10))

        # Executive Summary
        if meeting.summary:
            story.append(Paragraph("Executive Summary", h1_style))
            story.append(Paragraph(meeting.summary.quick_summary, body_style))

            if meeting.summary.detailed_summary:
                story.append(Paragraph("Detailed Summary", h1_style))
                for para in meeting.summary.detailed_summary.split("\n\n"):
                    if para.strip():
                        story.append(Paragraph(para.strip(), body_style))

            if meeting.summary.key_points:
                story.append(Paragraph("Key Discussion Points", h1_style))
                for kp in meeting.summary.key_points:
                    story.append(Paragraph(f"• {kp}", bullet_style))

            if meeting.summary.decisions:
                story.append(Paragraph("Decisions Made", h1_style))
                for d in meeting.summary.decisions:
                    story.append(Paragraph(f"✓ <b>{d}</b>", bullet_style))

            if meeting.summary.questions:
                story.append(Paragraph("Open Questions & Issues", h1_style))
                for q in meeting.summary.questions:
                    story.append(Paragraph(f"? {q}", bullet_style))

        # Action Items Table
        if meeting.action_items:
            story.append(Paragraph("Action Items", h1_style))
            table_data = [["Task", "Assignee", "Deadline", "Priority", "Status"]]
            for item in meeting.action_items:
                prio = item.priority.value if item.priority else "MEDIUM"
                stat = item.status.value if item.status else "PENDING"
                table_data.append([
                    Paragraph(item.task, body_style),
                    Paragraph(item.assignee, body_style),
                    Paragraph(item.deadline, body_style),
                    prio,
                    stat
                ])

            action_table = Table(table_data, colWidths=[200, 90, 80, 70, 70])
            action_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), primary_color),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(KeepTogether([action_table]))

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

